#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Markdown の議事録を .docx に変換する。

使い方:
    python make_docx.py <input.md> <output.docx>

環境変数 USE_PANDOC=1 かつ pandoc が入っていれば pandoc を使い、
それ以外は内蔵の簡易コンバータ(python-docx / 游ゴシック)で変換する。

内蔵コンバータの対応記法:
    見出し(# 〜 ####)、箇条書き(- / * / +、2 スペース字下げで入れ子)、
    番号付きリスト(1. )、表(| a | b |)、引用(>)、水平線(---)、
    **太字**、*斜体*、`コード`
"""
import os
import re
import shutil
import subprocess
import sys

INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+?`)")
HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
BULLET_RE = re.compile(r"^(\s*)([-*+])\s+(.*)$")
NUMBERED_RE = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
TABLE_SEP_RE = re.compile(r"^\s*\|?[\s:\-|]+\|?\s*$")


def convert_with_pandoc(src, dst):
    return subprocess.call(
        ["pandoc", src, "-f", "gfm", "-t", "docx", "-o", dst]) == 0


def add_runs(paragraph, text):
    """**太字** / *斜体* / `コード` を run に分解して段落へ追加する。"""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**"):
            run = paragraph.add_run(tok[2:-2])
            run.bold = True
        elif tok.startswith("`"):
            run = paragraph.add_run(tok[1:-1])
            run.font.name = "Courier New"
        else:
            run = paragraph.add_run(tok[1:-1])
            run.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def setup_fonts(doc):
    """日本語文書向けに主要スタイルを游ゴシックに揃える。"""
    from docx.oxml.ns import qn
    from docx.shared import Pt

    sizes = {
        "Normal": 10.5,
        "Title": 20,
        "Heading 1": 16,
        "Heading 2": 13,
        "Heading 3": 11.5,
        "Heading 4": 11,
    }
    for name, size in sizes.items():
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = "Yu Gothic"
        st.font.size = Pt(size)
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), "Yu Gothic")


def split_table_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def convert_builtin(src, dst):
    import docx

    doc = docx.Document()
    setup_fonts(doc)

    with open(src, encoding="utf-8") as f:
        lines = f.read().splitlines()

    i = 0
    used_title = False
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue

        # 水平線
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            i += 1
            continue

        # 見出し(最初の H1 は表題スタイル)
        m = HEADING_RE.match(stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and not used_title:
                p = doc.add_heading("", level=0)
                used_title = True
            else:
                p = doc.add_heading("", level=min(level, 4))
            add_runs(p, text)
            i += 1
            continue

        # 表
        if (stripped.startswith("|") and i + 1 < len(lines)
                and "-" in lines[i + 1]
                and TABLE_SEP_RE.match(lines[i + 1])):
            header = split_table_row(stripped)
            i += 2
            body = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                body.append(split_table_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for j, cell in enumerate(header):
                p = table.rows[0].cells[j].paragraphs[0]
                add_runs(p, cell)
                for r in p.runs:
                    r.bold = True
            for row in body:
                cells = table.add_row().cells
                for j in range(len(header)):
                    add_runs(cells[j].paragraphs[0],
                             row[j] if j < len(row) else "")
            continue

        # 箇条書き
        m = BULLET_RE.match(lines[i])
        if m:
            depth = min(len(m.group(1)) // 2, 2)
            style = "List Bullet" if depth == 0 else f"List Bullet {depth + 1}"
            p = doc.add_paragraph(style=style)
            add_runs(p, m.group(3))
            i += 1
            continue

        # 番号付きリスト
        m = NUMBERED_RE.match(lines[i])
        if m:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, m.group(2))
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            add_runs(p, stripped.lstrip("> ").strip())
            i += 1
            continue

        # 通常の段落(1 行 = 1 段落)
        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    doc.save(dst)


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    src, dst = sys.argv[1], sys.argv[2]

    use_pandoc = os.environ.get("USE_PANDOC", "0") == "1"
    if use_pandoc and shutil.which("pandoc"):
        if convert_with_pandoc(src, dst):
            print(f"[make_docx] pandoc で変換: {dst}")
            return
        print("[make_docx] pandoc 失敗、内蔵コンバータで再試行します")

    convert_builtin(src, dst)
    print(f"[make_docx] 変換完了: {dst}")


if __name__ == "__main__":
    main()
